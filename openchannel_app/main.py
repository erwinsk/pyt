# -*- coding: utf-8 -*-
"""
main.py
Aplikasi Desktop: Kalkulator Open Channel Flowmeter (Professional Edition)

Mendukung:
    - V-Notch Weir (Weir Segitiga)     : 2 metode
    - Rectangular Weir (Weir Persegi)  : 3 metode
    - Trapezoidal Weir (Weir Trapesium): 2 metode
    - Parshall Flume                   : 2 metode

Fitur:
    - Sketsa visual yang menyesuaikan tipe alat ukur, lengkap dengan label
      dimensi (H, kh, L, P, z, W) yang mengikuti nilai parameter secara live.
    - Menampilkan debit, laju alir massa (mass flow), estimasi kecepatan
      aliran & luas penampang basah.
    - Export riwayat perhitungan ke file Excel (.xlsx) atau CSV.

Menjalankan:
    pip install -r requirements.txt
    python main.py
"""

import sys
import os
import csv
import math
from datetime import datetime

from PyQt5.QtCore import Qt, QPointF, QTimer
from PyQt5.QtGui import (
    QPainter, QColor, QPen, QBrush, QIcon, QPixmap, QPolygonF, QPainterPath
)
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QFormLayout,
    QGroupBox, QComboBox, QDoubleSpinBox, QSpinBox, QPushButton, QLabel, QFrame,
    QSizePolicy, QMessageBox, QAction, QStatusBar, QFileDialog, QTabWidget,
    QCheckBox, QScrollArea
)

from flow_calc import TYPES
from modbus_reader import (
    ModbusRTUReader, list_serial_ports, decode_registers, register_count_for,
    PARITY_MAP, DATA_TYPE_OPTIONS, DATA_TYPE_KEY_MAP, WORD_ORDER_OPTIONS,
    WORD_ORDER_KEY_MAP, PYMODBUS_AVAILABLE, PYSERIAL_AVAILABLE, last_list_ports_error,
    PYMODBUS_VERSION,
)

try:
    import openpyxl
    from openpyxl.styles import Font as XLFont, PatternFill, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


UNIT_TO_M = {"mm": 0.001, "cm": 0.01, "m": 1.0}

COLOR_LINE = QColor("#33475b")
COLOR_PLATE = QColor("#8aa6bd")
COLOR_WATER = QColor("#a9d7f2")
COLOR_DIM = QColor("#c0392b")
COLOR_DIM2 = QColor("#b8860b")
COLOR_BG = QColor("#f7f9fc")


APP_STYLE = """
QMainWindow {
    background-color: #eef1f5;
}
QWidget {
    font-family: "Segoe UI", "Arial", sans-serif;
    font-size: 10pt;
    color: #1f2733;
}
QGroupBox {
    background-color: #ffffff;
    border: 1px solid #d3d9e2;
    border-radius: 8px;
    margin-top: 14px;
    padding: 12px;
    font-weight: 600;
    font-size: 10.5pt;
    color: #14477a;
}
QGroupBox::title {
    subcontrol-origin: margin;
    left: 12px;
    padding: 0 6px;
    color: #14477a;
}
QComboBox, QDoubleSpinBox {
    background-color: #fbfcfe;
    border: 1px solid #c3cbd6;
    border-radius: 5px;
    padding: 5px 8px;
    min-height: 22px;
}
QComboBox:focus, QDoubleSpinBox:focus {
    border: 1px solid #1f77c4;
}
QComboBox::drop-down {
    subcontrol-origin: padding;
    subcontrol-position: top right;
    border: none;
    border-left: 1px solid #c3cbd6;
    width: 22px;
}
QComboBox::down-arrow {
    image: none;
    width: 0px;
    height: 0px;
    border-left: 5px solid transparent;
    border-right: 5px solid transparent;
    border-top: 6px solid #33475b;
    margin-right: 7px;
}
QComboBox:disabled {
    background-color: #eef1f5;
    color: #9aa5b1;
}
QPushButton#calcBtn {
    background-color: #1f77c4;
    color: white;
    border-radius: 6px;
    padding: 10px 18px;
    font-weight: 600;
    font-size: 11pt;
}
QPushButton#calcBtn:hover { background-color: #1866aa; }
QPushButton#calcBtn:pressed { background-color: #124e85; }
QPushButton#resetBtn {
    background-color: #eceff3;
    color: #33475b;
    border: 1px solid #c3cbd6;
    border-radius: 6px;
    padding: 10px 14px;
    font-weight: 600;
}
QPushButton#resetBtn:hover { background-color: #dfe4ea; }
QPushButton#exportBtn {
    background-color: #1e8f5f;
    color: white;
    border-radius: 6px;
    padding: 8px 14px;
    font-weight: 600;
}
QPushButton#exportBtn:hover { background-color: #197a51; }
QPushButton#clearHistBtn {
    background-color: #eceff3;
    color: #8a3b3b;
    border: 1px solid #d9c3c3;
    border-radius: 6px;
    padding: 8px 12px;
    font-weight: 600;
}
QPushButton#clearHistBtn:hover { background-color: #f5e6e6; }
QLabel#headerTitle { color: #ffffff; font-size: 15pt; font-weight: 700; }
QLabel#headerSubtitle { color: #d7e8fb; font-size: 9pt; }
QFrame#headerBar { background-color: #14477a; }
QLabel#descLabel { color: #5b6b7f; font-style: italic; font-size: 9pt; padding: 4px 2px; }
QLabel#resultCaption { color: #5b6b7f; font-size: 9.5pt; font-weight: 600; }
QLabel#resultValue { color: #0b3d66; font-size: 13pt; font-weight: 700; }
QLabel#resultSub { color: #445261; font-size: 9pt; }
QLabel#historyLabel { color: #5b6b7f; font-size: 9pt; }
QStatusBar { background-color: #dde3ea; color: #445261; }
QFrame#hLine { color: #d3d9e2; }
"""


# =============================================================================
# SCHEMATIC WIDGET - sketsa visual per tipe alat ukur dengan label dimensi
# =============================================================================

class SchematicWidget(QWidget):
    """Menggambar sketsa (bukan gambar teknik presisi/skala) yang sesuai
    dengan tipe & metode yang dipilih, lengkap dengan label dimensi H, kh,
    L, P, z, W agar pengguna paham posisi tiap parameter."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.type_name = "V-Notch Weir (Weir Segitiga)"
        self.method_name = ""
        self.params = {}
        self.H_val = 150.0
        self.H_unit = "mm"
        self.setMinimumHeight(230)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

    def update_view(self, type_name, method_name, params, H_val, H_unit):
        self.type_name = type_name
        self.method_name = method_name
        self.params = params or {}
        self.H_val = H_val
        self.H_unit = H_unit
        self.update()

    # ---- helper drawing primitives -------------------------------------
    def _tick(self, p, x, y, horizontal_dim):
        """Tanda kecil (tick) tegak lurus garis dimensi."""
        x, y = int(round(x)), int(round(y))
        if horizontal_dim:
            p.drawLine(x, y - 4, x, y + 4)
        else:
            p.drawLine(x - 4, y, x + 4, y)

    def _draw_label(self, p, x, y, text, color=None, bg=True, line_height=14):
        """Menggambar label teks (boleh multi-baris, dipisah '\\n') dengan kotak
        latar putih semi-transparan di belakangnya.

        Dua masalah sekaligus diselesaikan di sini:
        1. QPainter.drawText(x, y, text) TIDAK merender karakter newline
           sebagai baris baru (hanya baris pertama yang tampak/rapi), jadi
           setiap baris digambar terpisah secara eksplisit.
        2. Label yang kebetulan jatuh di atas plat/air tidak lagi "menumpuk"
           tak terbaca, karena ada kotak latar yang memisahkannya secara
           visual dari gambar di belakangnya.
        """
        if not text:
            return
        color = color or COLOR_LINE
        lines = text.split("\n")
        fm = p.fontMetrics()
        if bg:
            width_fn = fm.horizontalAdvance if hasattr(fm, "horizontalAdvance") else fm.width
            max_w = max(width_fn(line) for line in lines)
            total_h = line_height * len(lines)
            p.save()
            p.setPen(Qt.NoPen)
            p.setBrush(QBrush(QColor(255, 255, 255, 220)))
            p.drawRoundedRect(int(x) - 3, int(y) - fm.ascent() - 3, int(max_w) + 6, int(total_h) + 6, 3, 3)
            p.restore()
        p.setPen(QPen(color, 1))
        for i, line in enumerate(lines):
            p.drawText(int(x), int(y) + i * line_height, line)

    def _dim_v(self, p, x, y_top, y_bottom, text, color=None, text_dx=8):
        """Garis dimensi VERTIKAL antara y_top dan y_bottom pada absis x,
        dengan tick di kedua ujung dan label teks di sisi kanan garis."""
        x = int(round(x))
        y_top, y_bottom = int(round(y_top)), int(round(y_bottom))
        color = color or COLOR_DIM
        pen = QPen(color, 1.4)
        p.setPen(pen)
        y1, y2 = sorted([y_top, y_bottom])
        p.drawLine(x, y1, x, y2)
        self._tick(p, x, y1, horizontal_dim=False)
        self._tick(p, x, y2, horizontal_dim=False)
        mid_y = (y1 + y2) / 2
        self._draw_label(p, x + text_dx, int(mid_y) + 4, text, color=color)

    def _dim_h(self, p, y, x_left, x_right, text, color=None, text_dy=-6):
        """Garis dimensi HORIZONTAL antara x_left dan x_right pada ordinat y."""
        y = int(round(y))
        x_left, x_right = int(round(x_left)), int(round(x_right))
        color = color or COLOR_DIM
        pen = QPen(color, 1.4)
        p.setPen(pen)
        x1, x2 = sorted([x_left, x_right])
        p.drawLine(x1, y, x2, y)
        self._tick(p, x1, y, horizontal_dim=True)
        self._tick(p, x2, y, horizontal_dim=True)
        if not text:
            return
        mid_x = (x1 + x2) / 2
        fm = p.fontMetrics()
        width_fn = fm.horizontalAdvance if hasattr(fm, "horizontalAdvance") else fm.width
        self._draw_label(p, int(mid_x) - width_fn(text) / 2, y + text_dy, text, color=color)

    def _fmt_H(self):
        return f"H = {self.H_val:g} {self.H_unit}"

    # ---- main paint ------------------------------------------------------
    def paintEvent(self, event):
        p = QPainter(self)
        p.setRenderHint(QPainter.Antialiasing)
        p.fillRect(self.rect(), COLOR_BG)

        if "V-Notch" in self.type_name:
            self._draw_vnotch(p)
        elif "Rectangular" in self.type_name:
            self._draw_rectangular(p)
        elif "Trapezoidal" in self.type_name:
            self._draw_trapezoidal(p)
        else:
            self._draw_parshall(p)

        p.end()

    # ---- V-NOTCH -----------------------------------------------------
    def _draw_vnotch(self, p):
        w, h = self.width(), self.height()
        left_x, right_x = 40, w - 132  # margin kanan lebih lega utk label H
        base_y = h - 25
        plate_top = 36  # ruang cukup untuk label 2 baris di atas plat
        apex_y = int(h * 0.60)
        notch_half_w = min(60, (right_x - left_x) * 0.32)
        cx = (left_x + right_x) / 2.0

        theta = self.params.get("theta", 90.0)
        Cd = self.params.get("Cd", None)
        has_kh = "kh" in self.params

        # dinding & dasar saluran
        pen_channel = QPen(COLOR_LINE, 2)
        p.setPen(pen_channel)
        p.drawLine(left_x, plate_top - 5, left_x, base_y)
        p.drawLine(right_x, plate_top - 5, right_x, base_y)
        p.drawLine(left_x, base_y, right_x, base_y)

        # water level (illustratif, offset tetap di atas apex)
        water_offset_px = 60
        water_y = apex_y - water_offset_px

        # plat weir (persegi penuh) dikurangi notch segitiga (path subtracted)
        plate_path = QPainterPath()
        plate_path.addRect(left_x, plate_top, right_x - left_x, base_y - plate_top)

        notch_path = QPainterPath()
        notch_path.moveTo(cx, apex_y)
        notch_path.lineTo(cx - notch_half_w, plate_top)
        notch_path.lineTo(cx + notch_half_w, plate_top)
        notch_path.closeSubpath()

        plate_minus_notch = plate_path.subtracted(notch_path)

        # air yang terlihat melalui notch (dari apex sampai water_y)
        water_path = QPainterPath()
        water_top_half_w = notch_half_w * (apex_y - water_y) / (apex_y - plate_top) if apex_y > plate_top else 0
        water_path.moveTo(cx, apex_y)
        water_path.lineTo(cx - water_top_half_w, water_y)
        water_path.lineTo(cx + water_top_half_w, water_y)
        water_path.closeSubpath()
        p.setPen(Qt.NoPen)
        p.setBrush(QBrush(COLOR_WATER))
        p.drawPath(water_path)

        # gambar plat (di atas air agar notch terlihat "memotong" plat)
        p.setBrush(QBrush(COLOR_PLATE))
        p.setPen(QPen(COLOR_LINE, 1.5))
        p.drawPath(plate_minus_notch)

        # garis muka air memanjang ke kanan channel
        p.setPen(QPen(QColor("#2f6fa8"), 1.5, Qt.DashLine))
        p.drawLine(left_x, water_y, right_x, water_y)

        # sudut theta di apex
        p.setPen(QPen(COLOR_DIM2, 1.3))
        p.drawArc(int(cx - 20), int(apex_y - 20), 40, 40, 60 * 16, 60 * 16)
        self._draw_label(p, cx + 10, apex_y - 22, f"\u03b8 = {theta:g}\u00b0", color=COLOR_DIM2)

        # dimensi H: dari apex (vertex notch, referensi H=0) sampai muka air
        self._dim_v(p, right_x + 20, apex_y, water_y, self._fmt_H(), color=COLOR_DIM)

        # kh (hanya bila metode memakai koreksi head) - garis & tick saja di sini;
        # penjelasan teksnya digabung ke legend kiri-bawah agar tidak menumpuk
        # dengan label lain di sekitar vertex/notch.
        kh_text = ""
        if has_kh:
            kh_val = self.params.get("kh", 0.0)
            kh_px = 14 if kh_val >= 0 else -14
            virt_y = apex_y - kh_px
            p.setPen(QPen(COLOR_DIM2, 1.2, Qt.DashLine))
            p.drawLine(int(cx - notch_half_w - 10), int(virt_y),
                       int(cx + notch_half_w + 10), int(virt_y))
            self._tick(p, cx - notch_half_w - 10, virt_y, horizontal_dim=False)
            kh_text = f"\nkh = {kh_val:g} mm (He = H+kh)"

        # label plat (di atas plat, di luar area gambar)
        self._draw_label(p, left_x, plate_top - 20,
                          "Plat V-Notch" + (f"\nCd = {Cd:g}" if Cd is not None else ""))

        # legend kiri-bawah: vertex + (opsional) kh, dikelompokkan jadi satu
        # blok teks supaya tidak berserakan/menumpuk di sekitar notch.
        p.setPen(QPen(COLOR_LINE, 1))
        p.drawEllipse(QPointF(cx, apex_y), 3, 3)
        self._draw_label(p, left_x, base_y - 34, "Vertex notch (H=0), diukur dari titik dasar V" + kh_text)

    # ---- RECTANGULAR ---------------------------------------------------
    def _draw_rectangular(self, p):
        w, h = self.width(), self.height()
        left_x, right_x = 40, w - 132  # margin kanan lebih lega utk label H
        base_y = h - 25
        plate_top = 36  # ruang cukup untuk label 2 baris di atas plat

        L = self.params.get("L", 0.5)
        P = self.params.get("P", None)  # hanya ada pada metode Rehbock
        Cd = self.params.get("Cd", None)

        crest_y = int(base_y - 60) if P is None else int(base_y - 45)
        crest_y = max(crest_y, plate_top + 45)
        notch_half_w = min(70, (right_x - left_x) * 0.35)
        cx = (left_x + right_x) / 2.0

        water_offset_px = 45
        water_y = crest_y - water_offset_px

        pen_channel = QPen(COLOR_LINE, 2)
        p.setPen(pen_channel)
        p.drawLine(left_x, plate_top - 5, left_x, base_y)
        p.drawLine(right_x, plate_top - 5, right_x, base_y)
        p.drawLine(left_x, base_y, right_x, base_y)

        # plat penuh dikurangi bukaan persegi (crest) di bagian atas
        plate_path = QPainterPath()
        plate_path.addRect(left_x, plate_top, right_x - left_x, base_y - plate_top)
        notch_path = QPainterPath()
        notch_path.addRect(cx - notch_half_w, plate_top, 2 * notch_half_w, crest_y - plate_top)
        plate_minus_notch = plate_path.subtracted(notch_path)

        # air yang terlihat melalui bukaan (dari crest sampai water_y)
        p.setPen(Qt.NoPen)
        p.setBrush(QBrush(COLOR_WATER))
        p.drawRect(int(cx - notch_half_w), int(water_y), int(2 * notch_half_w), int(crest_y - water_y))

        p.setBrush(QBrush(COLOR_PLATE))
        p.setPen(QPen(COLOR_LINE, 1.5))
        p.drawPath(plate_minus_notch)

        p.setPen(QPen(QColor("#2f6fa8"), 1.5, Qt.DashLine))
        p.drawLine(left_x, int(water_y), right_x, int(water_y))

        # dimensi H: dari crest sampai muka air
        self._dim_v(p, right_x + 20, crest_y, water_y, self._fmt_H(), color=COLOR_DIM)

        # dimensi P (jika Rehbock): dari dasar saluran sampai crest
        if P is not None:
            text_dx = max(-(left_x - 24 - 4), -70)  # jangan sampai x teks < 4px
            self._dim_v(p, left_x - 24, base_y, crest_y, f"P = {P:g} m",
                        color=COLOR_DIM2, text_dx=text_dx)

        # dimensi L: sepanjang bukaan crest
        self._dim_h(p, crest_y + 14, cx - notch_half_w, cx + notch_half_w, f"L = {L:g} m")

        self._draw_label(p, left_x, plate_top - 20,
                          "Weir Persegi" + (f"\nCd = {Cd:g}" if Cd is not None else ""))
        self._draw_label(p, left_x, crest_y - 10, "Mercu Weir (crest)")

    # ---- TRAPEZOIDAL -----------------------------------------------------
    def _draw_trapezoidal(self, p):
        w, h = self.width(), self.height()
        left_x, right_x = 40, w - 132  # margin kanan lebih lega utk label H
        base_y = h - 25
        plate_top = 36  # ruang cukup untuk label 2 baris di atas plat

        L = self.params.get("L", 0.5)
        z = self.params.get("z", 0.25)
        is_cipolletti = "z" not in self.params

        crest_y = int(base_y - 55)
        cx = (left_x + right_x) / 2.0
        bottom_half_w = min(45, (right_x - left_x) * 0.22)
        top_half_w = min(bottom_half_w + z * 55, (right_x - left_x) * 0.45)

        water_offset_px = 45
        water_y = crest_y - water_offset_px
        water_half_w_top = bottom_half_w + (top_half_w - bottom_half_w) * (
            (crest_y - water_y) / (crest_y - plate_top)
        )

        pen_channel = QPen(COLOR_LINE, 2)
        p.setPen(pen_channel)
        p.drawLine(left_x, plate_top - 5, left_x, base_y)
        p.drawLine(right_x, plate_top - 5, right_x, base_y)
        p.drawLine(left_x, base_y, right_x, base_y)

        # plat dikurangi bukaan trapesium
        plate_path = QPainterPath()
        plate_path.addRect(left_x, plate_top, right_x - left_x, base_y - plate_top)
        notch_path = QPainterPath()
        notch_path.moveTo(cx - bottom_half_w, crest_y)
        notch_path.lineTo(cx + bottom_half_w, crest_y)
        notch_path.lineTo(cx + top_half_w, plate_top)
        notch_path.lineTo(cx - top_half_w, plate_top)
        notch_path.closeSubpath()
        plate_minus_notch = plate_path.subtracted(notch_path)

        # air terlihat melalui bukaan trapesium (dari crest sampai water_y)
        water_path = QPainterPath()
        water_path.moveTo(cx - bottom_half_w, crest_y)
        water_path.lineTo(cx + bottom_half_w, crest_y)
        water_path.lineTo(cx + water_half_w_top, water_y)
        water_path.lineTo(cx - water_half_w_top, water_y)
        water_path.closeSubpath()
        p.setPen(Qt.NoPen)
        p.setBrush(QBrush(COLOR_WATER))
        p.drawPath(water_path)

        p.setBrush(QBrush(COLOR_PLATE))
        p.setPen(QPen(COLOR_LINE, 1.5))
        p.drawPath(plate_minus_notch)

        p.setPen(QPen(QColor("#2f6fa8"), 1.5, Qt.DashLine))
        p.drawLine(left_x, int(water_y), right_x, int(water_y))

        # dimensi H
        self._dim_v(p, right_x + 20, crest_y, water_y, self._fmt_H(), color=COLOR_DIM)

        # dimensi L (lebar dasar)
        self._dim_h(p, crest_y + 14, cx - bottom_half_w, cx + bottom_half_w, f"L = {L:g} m")

        # label kemiringan sisi z (diletakkan di luar bukaan, dengan latar
        # agar tidak menumpuk dengan plat/air di baliknya)
        mid_slope_x = (cx + bottom_half_w + cx + top_half_w) / 2.0
        mid_slope_y = (crest_y + plate_top) / 2.0
        if is_cipolletti:
            slope_text = "Kemiringan sisi\n1 : 4 (tetap - Cipolletti)"
        else:
            slope_text = f"Kemiringan sisi\n1 : {1 / z:.2f}  (z = {z:g})"
        self._draw_label(p, mid_slope_x + 4, mid_slope_y, slope_text, color=COLOR_DIM2)

        self._draw_label(p, left_x, plate_top - 20, "Weir Trapesium")
        self._draw_label(p, left_x, crest_y - 10, "Dasar bukaan trapesium")

    # ---- PARSHALL FLUME (denah / plan view) -------------------------------
    def _draw_parshall(self, p):
        w, h = self.width(), self.height()
        left_x, right_x = 35, w - 35
        mid_y = h / 2.0 - 5

        W_key = self.params.get("W_key", None)

        in_half = 55
        throat_half = 20
        out_half = 32

        conv_end_x = left_x + (right_x - left_x) * 0.42
        throat_end_x = left_x + (right_x - left_x) * 0.60

        pts = [
            QPointF(left_x, mid_y - in_half),
            QPointF(conv_end_x, mid_y - throat_half),
            QPointF(throat_end_x, mid_y - throat_half),
            QPointF(right_x, mid_y - out_half),
            QPointF(right_x, mid_y + out_half),
            QPointF(throat_end_x, mid_y + throat_half),
            QPointF(conv_end_x, mid_y + throat_half),
            QPointF(left_x, mid_y + in_half),
        ]
        poly = QPolygonF(pts)

        p.setPen(QPen(COLOR_LINE, 2))
        p.setBrush(QBrush(COLOR_WATER))
        p.drawPolygon(poly)

        # garis pemisah bagian konvergen/leher/divergen
        p.setPen(QPen(COLOR_LINE, 1, Qt.DashLine))
        p.drawLine(int(conv_end_x), int(mid_y - in_half - 10), int(conv_end_x), int(mid_y + in_half + 10))
        p.drawLine(int(throat_end_x), int(mid_y - in_half - 10), int(throat_end_x), int(mid_y + in_half + 10))

        self._draw_label(p, left_x + 5, mid_y - in_half - 26, "Bagian Konvergen\n(inlet)")
        self._draw_label(p, (conv_end_x + throat_end_x) / 2 - 22, mid_y - throat_half - 34, "Leher\n(Throat)")
        self._draw_label(p, throat_end_x + 8, mid_y - out_half - 26, "Bagian Divergen\n(outlet)")

        # dimensi lebar leher W
        w_label = f"W = {W_key}" if W_key else "W: koefisien kustom\n(tidak terkait lebar leher)"
        self._dim_h(p, mid_y - throat_half - 8, conv_end_x, throat_end_x, "", color=COLOR_DIM)
        self._draw_label(p, conv_end_x - 10, mid_y + throat_half + 45, w_label, color=COLOR_DIM)

        # titik pengukuran head Ha (kira-kira 2/3 dari bagian konvergen, sisi hulu)
        # label diletakkan DI BAWAH titik ukur (bukan di atas) supaya tidak
        # menumpuk dengan label "Leher (Throat)" yang ada di atasnya.
        gauge_x = left_x + (conv_end_x - left_x) * 0.65
        gauge_y = mid_y - (in_half + throat_half) / 2.0 * 0.55
        p.setPen(QPen(COLOR_DIM2, 1.6))
        p.setBrush(QBrush(COLOR_DIM2))
        p.drawEllipse(QPointF(gauge_x, gauge_y), 4, 4)
        self._draw_label(p, gauge_x + 8, gauge_y + 20, f"Titik ukur Ha\n({self._fmt_H()})", color=COLOR_DIM2)


# =============================================================================
# MAIN WINDOW
# =============================================================================

class MainWindow(QMainWindow):

    HISTORY_FIELDS = [
        "waktu", "sumber", "tipe", "metode", "parameter",
        "H_input", "H_satuan", "H_meter", "rho_kg_per_m3",
        "Q_m3_per_s", "Q_m3_per_jam", "Q_L_per_s", "Q_L_per_menit",
        "mdot_kg_per_s", "mdot_kg_per_jam", "mdot_ton_per_jam",
        "kecepatan_m_per_s", "luas_basah_m2", "catatan", "peringatan",
    ]

    SENSOR_LOG_FIELDS = [
        "waktu", "raw_registers", "nilai_decode_mentah",
        "H_konversi", "H_satuan", "H_meter",
        "tipe_alat_ukur", "metode",
        "Q_m3_per_s", "Q_L_per_s", "mdot_kg_per_s", "kecepatan_m_per_s",
    ]

    def __init__(self):
        super().__init__()
        self.setWindowTitle("Kalkulator Open Channel Flowmeter - Professional Edition")
        self.resize(1150, 760)
        self.setMinimumSize(1000, 700)
        self.setWindowIcon(self._make_icon())

        self.current_param_widgets = {}
        self.history = []
        self.last_result = None

        # --- state Modbus RTU ---
        self.modbus_reader = ModbusRTUReader()
        self.poll_timer = QTimer(self)
        self.poll_timer.timeout.connect(self.poll_once)
        self.is_polling = False
        self.log_file_path = None
        self.log_file_handle = None
        self.log_csv_writer = None
        self.consecutive_poll_errors = 0
        self.last_history_save_ts = None

        self._build_menu()
        self._build_ui()
        self._build_statusbar()

        self.type_combo.currentTextChanged.connect(self.on_type_changed)
        self.method_combo.currentTextChanged.connect(self.on_method_changed)
        self.calc_btn.clicked.connect(self.calculate)
        self.reset_btn.clicked.connect(self.reset_form)
        self.export_btn.clicked.connect(self.export_history)
        self.clear_hist_btn.clicked.connect(self.clear_history)
        self.H_input.valueChanged.connect(self.refresh_schematic)
        self.H_unit_combo.currentTextChanged.connect(self.refresh_schematic)

        # --- wiring tab Modbus ---
        self.mb_refresh_btn.clicked.connect(self.refresh_serial_ports)
        self.mb_connect_btn.clicked.connect(self.toggle_connection)
        self.mb_read_once_btn.clicked.connect(self.perform_single_read)
        self.mb_poll_toggle_btn.clicked.connect(self.toggle_polling)
        self.mb_log_checkbox.stateChanged.connect(self._on_log_checkbox_changed)
        self.mb_log_path_btn.clicked.connect(self.choose_log_file)

        self.type_combo.addItems(list(TYPES.keys()))
        self.on_type_changed(self.type_combo.currentText())
        self._update_history_label()
        self.refresh_serial_ports()
        self._update_connection_status(False)

    # ------------------------------------------------------------------
    def _make_icon(self):
        pix = QPixmap(64, 64)
        pix.fill(Qt.transparent)
        p = QPainter(pix)
        p.setRenderHint(QPainter.Antialiasing)
        p.setBrush(QColor("#1f77c4"))
        p.setPen(Qt.NoPen)
        p.drawEllipse(2, 2, 60, 60)
        p.setPen(QPen(QColor("#ffffff"), 4))
        p.drawArc(14, 30, 36, 20, 0, 180 * 16)
        p.drawArc(14, 18, 36, 20, 0, 180 * 16)
        p.end()
        return QIcon(pix)

    def _build_menu(self):
        menubar = self.menuBar()
        file_menu = menubar.addMenu("&File")

        export_action = QAction("Export Riwayat ke Excel/CSV...", self)
        export_action.triggered.connect(self.export_history)
        file_menu.addAction(export_action)

        clear_action = QAction("Hapus Riwayat", self)
        clear_action.triggered.connect(self.clear_history)
        file_menu.addAction(clear_action)

        file_menu.addSeparator()
        exit_action = QAction("Keluar", self)
        exit_action.triggered.connect(self.close)
        file_menu.addAction(exit_action)

        help_menu = menubar.addMenu("&Bantuan")
        about_action = QAction("Tentang Aplikasi", self)
        about_action.triggered.connect(self.show_about)
        help_menu.addAction(about_action)

    def _build_statusbar(self):
        sb = QStatusBar()
        self.setStatusBar(sb)
        self.status_label = QLabel("Siap. Pilih tipe alat ukur, isi parameter, lalu klik Hitung.")
        sb.addWidget(self.status_label)

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        outer = QVBoxLayout(central)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setSpacing(0)

        header = QFrame()
        header.setObjectName("headerBar")
        header.setMinimumHeight(64)
        h_layout = QVBoxLayout(header)
        h_layout.setContentsMargins(20, 8, 20, 8)
        h_layout.setSpacing(0)
        title = QLabel("Kalkulator Open Channel Flowmeter")
        title.setObjectName("headerTitle")
        subtitle = QLabel("V-Notch Weir \u2022 Rectangular Weir \u2022 Trapezoidal Weir \u2022 Parshall Flume")
        subtitle.setObjectName("headerSubtitle")
        h_layout.addWidget(title)
        h_layout.addWidget(subtitle)
        outer.addWidget(header)

        tabs = QTabWidget()
        outer.addWidget(tabs)

        calc_inner = QWidget()
        body_layout = QHBoxLayout(calc_inner)
        body_layout.setContentsMargins(16, 16, 16, 16)
        body_layout.setSpacing(16)
        body_layout.addWidget(self._build_input_panel(), stretch=5)
        body_layout.addWidget(self._build_result_panel(), stretch=4)

        calc_scroll = QScrollArea()
        calc_scroll.setWidgetResizable(True)
        calc_scroll.setWidget(calc_inner)
        calc_scroll.setFrameShape(QFrame.NoFrame)
        tabs.addTab(calc_scroll, "Kalkulasi")

        modbus_tab = self._build_modbus_tab()
        tabs.addTab(modbus_tab, "Modbus RTU - Pembacaan Sensor")
        self._modbus_tab_index = tabs.indexOf(modbus_tab)
        tabs.currentChanged.connect(self._on_tab_changed)

    def _build_input_panel(self):
        group = QGroupBox("Parameter Input")
        layout = QVBoxLayout(group)
        layout.setSpacing(10)

        form_top = QFormLayout()
        form_top.setSpacing(8)

        self.type_combo = QComboBox()
        form_top.addRow("Tipe Alat Ukur:", self.type_combo)

        self.method_combo = QComboBox()
        form_top.addRow("Metode Perhitungan:", self.method_combo)

        layout.addLayout(form_top)

        self.desc_label = QLabel("")
        self.desc_label.setObjectName("descLabel")
        self.desc_label.setWordWrap(True)
        layout.addWidget(self.desc_label)

        layout.addWidget(self._hline())

        self.schematic = SchematicWidget()
        layout.addWidget(self.schematic)

        layout.addWidget(self._hline())

        self.param_form = QFormLayout()
        self.param_form.setSpacing(8)
        layout.addLayout(self.param_form)

        layout.addWidget(self._hline())

        form_bottom = QFormLayout()
        form_bottom.setSpacing(8)

        h_row = QHBoxLayout()
        self.H_input = QDoubleSpinBox()
        self.H_input.setRange(0.001, 100000.0)
        self.H_input.setDecimals(3)
        self.H_input.setSingleStep(1.0)
        self.H_input.setValue(150.0)
        self.H_unit_combo = QComboBox()
        self.H_unit_combo.addItems(["mm", "cm", "m"])
        self.H_unit_combo.setCurrentText("mm")
        h_row.addWidget(self.H_input, stretch=3)
        h_row.addWidget(self.H_unit_combo, stretch=1)
        form_bottom.addRow("Ketinggian Air / Head (H):", h_row)

        self.rho_input = QDoubleSpinBox()
        self.rho_input.setRange(1.0, 5000.0)
        self.rho_input.setDecimals(1)
        self.rho_input.setSingleStep(1.0)
        self.rho_input.setValue(1000.0)
        self.rho_input.setSuffix(" kg/m\u00b3")
        form_bottom.addRow("Densitas Fluida (\u03c1):", self.rho_input)

        layout.addLayout(form_bottom)

        btn_row = QHBoxLayout()
        self.calc_btn = QPushButton("Hitung")
        self.calc_btn.setObjectName("calcBtn")
        self.reset_btn = QPushButton("Reset")
        self.reset_btn.setObjectName("resetBtn")
        btn_row.addWidget(self.calc_btn, stretch=3)
        btn_row.addWidget(self.reset_btn, stretch=1)
        layout.addLayout(btn_row)

        layout.addStretch(1)
        return group

    def _build_result_panel(self):
        group = QGroupBox("Hasil Perhitungan")
        layout = QVBoxLayout(group)
        layout.setSpacing(14)

        def make_result_block(caption):
            cap = QLabel(caption)
            cap.setObjectName("resultCaption")
            val = QLabel("-")
            val.setObjectName("resultValue")
            val.setWordWrap(True)
            sub = QLabel("")
            sub.setObjectName("resultSub")
            sub.setWordWrap(True)
            box = QVBoxLayout()
            box.setSpacing(2)
            box.addWidget(cap)
            box.addWidget(val)
            box.addWidget(sub)
            return box, val, sub

        box_q, self.val_Q, self.sub_Q = make_result_block("DEBIT VOLUMETRIK (Q)")
        layout.addLayout(box_q)
        layout.addWidget(self._hline())

        box_m, self.val_M, self.sub_M = make_result_block("LAJU ALIR MASSA (\u1e41)")
        layout.addLayout(box_m)
        layout.addWidget(self._hline())

        box_v, self.val_V, self.sub_V = make_result_block("KECEPATAN ALIRAN (v) - estimasi")
        layout.addLayout(box_v)
        layout.addWidget(self._hline())

        box_a, self.val_A, self.sub_A = make_result_block("LUAS PENAMPANG BASAH (A) - estimasi")
        layout.addLayout(box_a)
        layout.addWidget(self._hline())

        note_caption = QLabel("CATATAN & PERINGATAN")
        note_caption.setObjectName("resultCaption")
        layout.addWidget(note_caption)
        self.notes_label = QLabel("Belum ada perhitungan.")
        self.notes_label.setWordWrap(True)
        self.notes_label.setObjectName("resultSub")
        layout.addWidget(self.notes_label)

        layout.addWidget(self._hline())

        # ---- Export & riwayat ----
        export_caption = QLabel("EXPORT RIWAYAT PERHITUNGAN")
        export_caption.setObjectName("resultCaption")
        layout.addWidget(export_caption)

        self.history_label = QLabel("Riwayat: 0 perhitungan tersimpan di sesi ini.")
        self.history_label.setObjectName("historyLabel")
        layout.addWidget(self.history_label)

        export_row = QHBoxLayout()
        self.export_btn = QPushButton("Export ke Excel/CSV")
        self.export_btn.setObjectName("exportBtn")
        self.clear_hist_btn = QPushButton("Hapus Riwayat")
        self.clear_hist_btn.setObjectName("clearHistBtn")
        export_row.addWidget(self.export_btn, stretch=3)
        export_row.addWidget(self.clear_hist_btn, stretch=2)
        layout.addLayout(export_row)

        layout.addStretch(1)

        disclaimer = QLabel(
            "Catatan: Kecepatan & luas penampang adalah ESTIMASI pada seksi kontrol "
            "berdasarkan geometri sederhana, bukan pengukuran langsung. Sketsa bersifat "
            "ilustratif (tidak selalu skalatis). Untuk pengukuran resmi/legal, verifikasi "
            "koefisien terhadap standar yang berlaku."
        )
        disclaimer.setWordWrap(True)
        disclaimer.setStyleSheet("color:#8a94a3; font-size:8.5pt;")
        layout.addWidget(disclaimer)

        return group

    def _hline(self):
        line = QFrame()
        line.setObjectName("hLine")
        line.setFrameShape(QFrame.HLine)
        line.setFrameShadow(QFrame.Sunken)
        return line

    # ==================================================================
    # TAB MODBUS RTU - Pembacaan otomatis ketinggian air dari sensor
    # ==================================================================
    def _build_modbus_tab(self):
        tab = QWidget()
        outer = QHBoxLayout(tab)
        outer.setContentsMargins(16, 16, 16, 16)
        outer.setSpacing(16)

        left_col = QVBoxLayout()
        left_col.setSpacing(14)
        right_col = QVBoxLayout()
        right_col.setSpacing(14)

        if not PYMODBUS_AVAILABLE or not PYSERIAL_AVAILABLE:
            warn = QLabel(
                "\u26a0 Modul 'pymodbus' dan/atau 'pyserial' belum terpasang. "
                "Fitur pembacaan sensor akan menampilkan pesan error saat digunakan.\n"
                "Install dengan: pip install pymodbus pyserial"
            )
            warn.setWordWrap(True)
            warn.setStyleSheet(
                "background-color:#fdf2e0; color:#a6410c; border:1px solid #f0c896;"
                "border-radius:6px; padding:8px; font-size:9pt;"
            )
            left_col.addWidget(warn)

        if PYMODBUS_AVAILABLE:
            version_text = f"pymodbus terpasang: v{PYMODBUS_VERSION}" if PYMODBUS_VERSION else \
                "pymodbus terpasang (versi tidak terdeteksi)"
            version_label = QLabel(version_text)
            version_label.setStyleSheet("color:#6b7280; font-size:8pt;")
            left_col.addWidget(version_label)

        left_col.addWidget(self._build_connection_group())
        left_col.addWidget(self._build_mapping_group())
        left_col.addStretch(1)

        right_col.addWidget(self._build_control_group())
        right_col.addWidget(self._build_live_group())
        right_col.addStretch(1)

        left_wrap = QWidget()
        left_wrap.setLayout(left_col)
        right_wrap = QWidget()
        right_wrap.setLayout(right_col)

        outer.addWidget(left_wrap, stretch=5)
        outer.addWidget(right_wrap, stretch=5)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setWidget(tab)
        scroll.setFrameShape(QFrame.NoFrame)
        return scroll

    def _build_connection_group(self):
        group = QGroupBox("Koneksi Serial (RS-485 / Modbus RTU)")
        form = QFormLayout(group)
        form.setSpacing(8)

        port_row = QHBoxLayout()
        self.mb_port_combo = QComboBox()
        self.mb_port_combo.setEditable(True)
        self.mb_port_combo.setInsertPolicy(QComboBox.NoInsert)
        self.mb_port_combo.setMinimumWidth(140)
        self.mb_port_combo.setToolTip(
            "Pilih port dari daftar (klik panah di kanan), atau ketik manual "
            "bila port tidak terdeteksi otomatis."
        )
        self.mb_refresh_btn = QPushButton("\u27f3 Refresh")  # ⟳ Refresh
        self.mb_refresh_btn.setObjectName("resetBtn")
        self.mb_refresh_btn.setToolTip("Pindai ulang daftar port serial yang tersedia")
        port_row.addWidget(self.mb_port_combo, stretch=3)
        port_row.addWidget(self.mb_refresh_btn, stretch=1)
        form.addRow("Port Serial:", port_row)

        self.mb_baud_combo = QComboBox()
        self.mb_baud_combo.addItems(["1200", "2400", "4800", "9600", "19200", "38400", "57600", "115200"])
        self.mb_baud_combo.setCurrentText("9600")
        form.addRow("Baudrate:", self.mb_baud_combo)

        self.mb_parity_combo = QComboBox()
        self.mb_parity_combo.addItems(list(PARITY_MAP.keys()))
        form.addRow("Parity:", self.mb_parity_combo)

        self.mb_databits_combo = QComboBox()
        self.mb_databits_combo.addItems(["7", "8"])
        self.mb_databits_combo.setCurrentText("8")
        form.addRow("Data Bits:", self.mb_databits_combo)

        self.mb_stopbits_combo = QComboBox()
        self.mb_stopbits_combo.addItems(["1", "2"])
        form.addRow("Stop Bits:", self.mb_stopbits_combo)

        self.mb_slaveid_input = QSpinBox()
        self.mb_slaveid_input.setRange(1, 247)
        self.mb_slaveid_input.setValue(1)
        form.addRow("Slave ID (Unit ID):", self.mb_slaveid_input)

        self.mb_timeout_input = QDoubleSpinBox()
        self.mb_timeout_input.setRange(0.1, 10.0)
        self.mb_timeout_input.setDecimals(1)
        self.mb_timeout_input.setSingleStep(0.1)
        self.mb_timeout_input.setValue(1.0)
        self.mb_timeout_input.setSuffix(" s")
        form.addRow("Timeout:", self.mb_timeout_input)

        status_row = QHBoxLayout()
        self.mb_status_dot = QFrame()
        self.mb_status_dot.setFixedSize(14, 14)
        self.mb_status_label = QLabel("Terputus")
        self.mb_connect_btn = QPushButton("Sambungkan")
        self.mb_connect_btn.setObjectName("calcBtn")
        status_row.addWidget(self.mb_status_dot)
        status_row.addWidget(self.mb_status_label, stretch=1)
        status_row.addWidget(self.mb_connect_btn, stretch=2)
        form.addRow("Status:", status_row)

        return group

    def _build_mapping_group(self):
        group = QGroupBox("Pemetaan Register (Register Mapping)")
        form = QFormLayout(group)
        form.setSpacing(8)

        self.mb_func_combo = QComboBox()
        self.mb_func_combo.addItems(["Read Holding Registers (FC03)", "Read Input Registers (FC04)"])
        form.addRow("Fungsi Baca:", self.mb_func_combo)

        self.mb_address_input = QSpinBox()
        self.mb_address_input.setRange(0, 65535)
        self.mb_address_input.setValue(0)
        form.addRow("Alamat Register (0-based):", self.mb_address_input)

        self.mb_datatype_combo = QComboBox()
        self.mb_datatype_combo.addItems(DATA_TYPE_OPTIONS)
        self.mb_datatype_combo.setCurrentIndex(1)  # int16 (umum utk banyak transmitter)
        form.addRow("Tipe Data:", self.mb_datatype_combo)

        self.mb_wordorder_combo = QComboBox()
        self.mb_wordorder_combo.addItems(WORD_ORDER_OPTIONS)
        form.addRow("Urutan Word/Byte (32-bit):", self.mb_wordorder_combo)

        self.mb_scale_input = QDoubleSpinBox()
        self.mb_scale_input.setRange(-1000000.0, 1000000.0)
        self.mb_scale_input.setDecimals(6)
        self.mb_scale_input.setValue(1.0)
        form.addRow("Faktor Skala (Multiplier):", self.mb_scale_input)

        self.mb_offset_input = QDoubleSpinBox()
        self.mb_offset_input.setRange(-1000000.0, 1000000.0)
        self.mb_offset_input.setDecimals(6)
        self.mb_offset_input.setValue(0.0)
        form.addRow("Offset (Penambah):", self.mb_offset_input)

        self.mb_unit_combo = QComboBox()
        self.mb_unit_combo.addItems(["mm", "cm", "m"])
        form.addRow("Satuan Hasil (setelah skala):", self.mb_unit_combo)

        hint = QLabel(
            "Nilai H = (register_terbaca \u00d7 Faktor Skala) + Offset, dalam satuan di atas.\n"
            "Contoh umum: sensor keluaran 0-20000 = 0-20.000 mm \u2192 tipe uint16/int16, "
            "skala = 1, satuan = mm."
        )
        hint.setWordWrap(True)
        hint.setObjectName("descLabel")
        form.addRow(hint)

        return group

    def _build_control_group(self):
        group = QGroupBox("Kontrol Pembacaan")
        layout = QVBoxLayout(group)
        layout.setSpacing(8)

        self.mb_read_once_btn = QPushButton("Baca Sensor Sekali")
        self.mb_read_once_btn.setObjectName("exportBtn")
        layout.addWidget(self.mb_read_once_btn)

        self.mb_raw_label = QLabel("Register mentah: -\nNilai decode: -")
        self.mb_raw_label.setObjectName("resultSub")
        self.mb_raw_label.setWordWrap(True)
        layout.addWidget(self.mb_raw_label)

        layout.addWidget(self._hline())

        form = QFormLayout()
        form.setSpacing(8)
        self.mb_interval_input = QDoubleSpinBox()
        self.mb_interval_input.setRange(0.2, 3600.0)
        self.mb_interval_input.setDecimals(1)
        self.mb_interval_input.setValue(2.0)
        self.mb_interval_input.setSuffix(" s")
        form.addRow("Interval Polling:", self.mb_interval_input)
        layout.addLayout(form)

        autocalc_note = QLabel(
            "Setiap pembacaan (sekali maupun polling) otomatis mengisi H di tab "
            "Kalkulasi DAN langsung menghitung ulang debit/mass flow memakai "
            "Tipe & Metode yang sedang aktif di sana \u2014 pembacaan sensor di sini "
            "bukan fitur pemantauan ketinggian air yang berdiri sendiri."
        )
        autocalc_note.setWordWrap(True)
        autocalc_note.setObjectName("descLabel")
        layout.addWidget(autocalc_note)

        self.mb_log_checkbox = QCheckBox("Simpan Log ke CSV Setiap Polling")
        layout.addWidget(self.mb_log_checkbox)

        log_row = QHBoxLayout()
        self.mb_log_path_label = QLabel("(belum dipilih)")
        self.mb_log_path_label.setObjectName("resultSub")
        self.mb_log_path_label.setWordWrap(True)
        self.mb_log_path_btn = QPushButton("Pilih File Log...")
        self.mb_log_path_btn.setObjectName("resetBtn")
        log_row.addWidget(self.mb_log_path_label, stretch=3)
        log_row.addWidget(self.mb_log_path_btn, stretch=1)
        layout.addLayout(log_row)

        layout.addWidget(self._hline())

        self.mb_history_checkbox = QCheckBox("Simpan Sampel Polling ke Riwayat Kalkulasi (untuk Export)")
        layout.addWidget(self.mb_history_checkbox)

        history_form = QFormLayout()
        history_form.setSpacing(8)
        self.mb_history_interval_input = QDoubleSpinBox()
        self.mb_history_interval_input.setRange(1.0, 3600.0)
        self.mb_history_interval_input.setDecimals(1)
        self.mb_history_interval_input.setValue(10.0)
        self.mb_history_interval_input.setSuffix(" s")
        self.mb_history_interval_input.setEnabled(False)
        history_form.addRow("Interval Simpan ke Riwayat:", self.mb_history_interval_input)
        layout.addLayout(history_form)

        self.mb_history_checkbox.toggled.connect(self.mb_history_interval_input.setEnabled)

        history_note = QLabel(
            "Terpisah dari log CSV mentah di atas: opsi ini menambahkan baris "
            "hasil polling (H, Q, mass flow, dst.) langsung ke Riwayat tab "
            "Kalkulasi, sehingga ikut terekspor bersama hasil perhitungan "
            "manual saat menekan \u201cExport Riwayat\u201d. Karena polling bisa "
            "berjalan sangat cepat, hanya satu sampel yang disimpan per "
            "interval di atas (bukan setiap siklus polling)."
        )
        history_note.setWordWrap(True)
        history_note.setObjectName("descLabel")
        layout.addWidget(history_note)

        layout.addWidget(self._hline())

        self.mb_poll_toggle_btn = QPushButton("Mulai Polling")
        self.mb_poll_toggle_btn.setObjectName("calcBtn")
        layout.addWidget(self.mb_poll_toggle_btn)

        return group

    def _build_live_group(self):
        group = QGroupBox("Pembacaan Live")
        layout = QVBoxLayout(group)
        layout.setSpacing(4)

        cap = QLabel("KETINGGIAN AIR TERKINI (H)")
        cap.setObjectName("resultCaption")
        layout.addWidget(cap)

        self.mb_live_h_label = QLabel("-")
        self.mb_live_h_label.setObjectName("resultValue")
        layout.addWidget(self.mb_live_h_label)

        self.mb_live_time_label = QLabel("Belum ada pembacaan.")
        self.mb_live_time_label.setObjectName("resultSub")
        layout.addWidget(self.mb_live_time_label)

        note = QLabel(
            "Nilai H di sini otomatis mengisi field \u201cKetinggian Air / Head (H)\u201d "
            "di tab Kalkulasi, dan debit/mass flow ikut dihitung ulang secara "
            "otomatis di sana \u2014 baik saat \u201cBaca Sensor Sekali\u201d maupun "
            "selama polling berjalan."
        )
        note.setWordWrap(True)
        note.setObjectName("descLabel")
        layout.addWidget(note)

        return group

    # ---- logika koneksi & pembacaan Modbus --------------------------------
    def _on_tab_changed(self, index):
        if index == self._modbus_tab_index:
            self.refresh_serial_ports()

    def refresh_serial_ports(self):
        current = self.mb_port_combo.currentText()
        self.mb_port_combo.clear()
        ports = list_serial_ports()
        self.mb_port_combo.addItems(ports)
        if current:
            self.mb_port_combo.setCurrentText(current)

        if not PYSERIAL_AVAILABLE:
            self.status_label.setText(
                "Modul 'pyserial' belum terpasang — install dengan: pip install pyserial"
            )
        elif not ports:
            err = last_list_ports_error()
            if err:
                self.status_label.setText(f"Gagal mendeteksi port serial: {err}")
            else:
                self.status_label.setText(
                    "Tidak ada port serial terdeteksi. Pastikan perangkat RS-485/USB sudah tersambung."
                )
        else:
            self.status_label.setText(f"{len(ports)} port serial ditemukan.")

    def _update_connection_status(self, connected):
        color = "#1e8f5f" if connected else "#c0392b"
        self.mb_status_dot.setStyleSheet(f"background-color:{color}; border-radius:7px;")
        self.mb_status_label.setText("Terhubung" if connected else "Terputus")
        self.mb_connect_btn.setText("Putuskan Koneksi" if connected else "Sambungkan")

    def toggle_connection(self):
        if not self.modbus_reader.connected:
            port = self.mb_port_combo.currentText().strip()
            baud = int(self.mb_baud_combo.currentText())
            parity = PARITY_MAP[self.mb_parity_combo.currentText()]
            bytesize = int(self.mb_databits_combo.currentText())
            stopbits = int(self.mb_stopbits_combo.currentText())
            timeout = self.mb_timeout_input.value()

            ok = self.modbus_reader.connect(port, baud, parity, bytesize, stopbits, timeout)
            if ok:
                self._update_connection_status(True)
                self.status_label.setText(f"Terhubung ke {port} ({baud} bps).")
            else:
                QMessageBox.warning(self, "Koneksi Gagal",
                                     self.modbus_reader.last_error or "Gagal terhubung ke perangkat.")
        else:
            if self.is_polling:
                self.toggle_polling()
            self.modbus_reader.disconnect()
            self._update_connection_status(False)
            self.status_label.setText("Koneksi Modbus diputus.")

    def _current_mapping(self):
        data_type_label = self.mb_datatype_combo.currentText()
        word_order_label = self.mb_wordorder_combo.currentText()
        function_label = self.mb_func_combo.currentText()
        return dict(
            function_code="input" if "FC04" in function_label else "holding",
            address=self.mb_address_input.value(),
            data_type=DATA_TYPE_KEY_MAP[data_type_label],
            word_order=WORD_ORDER_KEY_MAP[word_order_label],
            scale=self.mb_scale_input.value(),
            offset=self.mb_offset_input.value(),
            unit=self.mb_unit_combo.currentText(),
            slave_id=self.mb_slaveid_input.value(),
        )

    def _read_and_decode(self):
        """Membaca & mendekode 1x sesuai pemetaan aktif.
        Mengembalikan (registers, raw_value, converted_value, unit) atau None jika gagal
        (pesan error tersimpan di self.modbus_reader.last_error)."""
        m = self._current_mapping()
        count = register_count_for(m["data_type"])
        registers = self.modbus_reader.read_raw_registers(
            m["slave_id"], m["address"], count, m["function_code"]
        )
        if registers is None:
            return None
        try:
            raw_value = decode_registers(registers, m["data_type"], m["word_order"])
        except Exception as e:
            self.modbus_reader.last_error = str(e)
            return None
        converted = raw_value * m["scale"] + m["offset"]
        return registers, raw_value, converted, m["unit"]

    def perform_single_read(self):
        if not self.modbus_reader.connected:
            QMessageBox.warning(self, "Belum Terhubung", "Sambungkan ke perangkat Modbus terlebih dahulu.")
            return
        result = self._read_and_decode()
        if result is None:
            QMessageBox.warning(self, "Gagal Membaca Sensor",
                                 self.modbus_reader.last_error or "Gagal membaca register.")
            return
        registers, raw_value, converted, unit = result
        self.mb_raw_label.setText(f"Register mentah: {registers}\nNilai decode: {raw_value:g}")
        self.H_input.setValue(converted)
        self.H_unit_combo.setCurrentText(unit)
        self.mb_live_h_label.setText(f"{converted:,.3f} {unit}")
        self.mb_live_time_label.setText(f"Terakhir dibaca: {datetime.now().strftime('%H:%M:%S')} (baca sekali)")
        # Tujuan utama pembacaan Modbus adalah mengisi H di tab Kalkulasi
        # secara otomatis (bukan sekadar menampilkan nilai sensor berdiri
        # sendiri), jadi langsung picu perhitungan debit dengan Tipe/Metode
        # yang sedang aktif di tab Kalkulasi.
        self.calculate(sumber="Sensor Modbus (Baca Sekali)")
        self.status_label.setText(f"Pembacaan sensor berhasil: H = {converted:,.3f} {unit} - debit dihitung otomatis.")

    def _on_log_checkbox_changed(self):
        if self.mb_log_checkbox.isChecked() and not self.log_file_path:
            self.choose_log_file()
            if not self.log_file_path:
                self.mb_log_checkbox.setChecked(False)

    def choose_log_file(self):
        default_name = f"log_sensor_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        fname, _ = QFileDialog.getSaveFileName(self, "Pilih File Log CSV", default_name, "CSV File (*.csv)")
        if not fname:
            return
        if not fname.lower().endswith(".csv"):
            fname += ".csv"
        self.log_file_path = fname
        self.mb_log_path_label.setText(fname)

    def _open_log_file(self):
        is_new = (not os.path.exists(self.log_file_path)) or os.path.getsize(self.log_file_path) == 0
        self.log_file_handle = open(self.log_file_path, "a", newline="", encoding="utf-8-sig")
        self.log_csv_writer = csv.DictWriter(self.log_file_handle, fieldnames=self.SENSOR_LOG_FIELDS)
        if is_new:
            self.log_csv_writer.writeheader()
            self.log_file_handle.flush()

    def _write_log_row(self, row):
        if self.log_csv_writer is None:
            return
        try:
            self.log_csv_writer.writerow(row)
            self.log_file_handle.flush()
        except Exception as e:
            self.status_label.setText(f"Gagal menulis log: {e}")

    def _close_log_file(self):
        if self.log_file_handle is not None:
            try:
                self.log_file_handle.close()
            except Exception:
                pass
        self.log_file_handle = None
        self.log_csv_writer = None

    def toggle_polling(self):
        if not self.is_polling:
            if not self.modbus_reader.connected:
                QMessageBox.warning(self, "Belum Terhubung", "Sambungkan ke perangkat Modbus terlebih dahulu.")
                return
            if self.mb_log_checkbox.isChecked():
                if not self.log_file_path:
                    self.choose_log_file()
                if not self.log_file_path:
                    QMessageBox.warning(self, "File Log Belum Dipilih",
                                         "Pilih file log CSV terlebih dahulu, atau nonaktifkan opsi logging.")
                    return
                try:
                    self._open_log_file()
                except Exception as e:
                    QMessageBox.critical(self, "Gagal Membuka File Log", str(e))
                    return

            self.consecutive_poll_errors = 0
            self.last_history_save_ts = None
            interval_ms = max(int(self.mb_interval_input.value() * 1000), 200)
            self.poll_timer.start(interval_ms)
            self.is_polling = True
            self.mb_poll_toggle_btn.setText("Hentikan Polling")
            self._restyle_button(self.mb_poll_toggle_btn, "clearHistBtn")
            self.status_label.setText("Polling Modbus dimulai.")
        else:
            self.poll_timer.stop()
            self.is_polling = False
            self.mb_poll_toggle_btn.setText("Mulai Polling")
            self._restyle_button(self.mb_poll_toggle_btn, "calcBtn")
            self._close_log_file()
            self.status_label.setText("Polling Modbus dihentikan.")

    def _restyle_button(self, button, object_name):
        """Mengganti objectName tombol lalu memaksa Qt menerapkan ulang QSS
        (setObjectName saja tidak otomatis memicu re-polish stylesheet)."""
        button.setObjectName(object_name)
        button.style().unpolish(button)
        button.style().polish(button)

    def _compute_flow_quiet(self, H_m):
        """Sama seperti calculate(), tapi tanpa dialog error dan tanpa
        menambah riwayat kalkulasi manual - dipakai saat auto-calc polling."""
        try:
            type_name = self.type_combo.currentText()
            method_name = self.method_combo.currentText()
            entry = TYPES[type_name]["methods"][method_name]
            func = entry["func"]
            params = self._get_param_values()
            rho = self.rho_input.value()
            Q, area, notes, warnings = func(H_m, params, rho)
            if Q is None or math.isnan(Q) or math.isinf(Q) or Q < 0:
                return None
            return Q, area, notes, warnings, rho
        except Exception:
            return None

    def poll_once(self):
        result = self._read_and_decode()
        if result is None:
            self.consecutive_poll_errors += 1
            self.status_label.setText(
                f"Polling gagal ({self.consecutive_poll_errors}x): "
                f"{self.modbus_reader.last_error or 'error tidak diketahui'}"
            )
            if self.consecutive_poll_errors >= 5:
                self.toggle_polling()
                QMessageBox.warning(self, "Polling Dihentikan",
                                     "Polling dihentikan otomatis setelah 5 kali gagal berturut-turut. "
                                     "Periksa koneksi & pemetaan register.")
            return

        self.consecutive_poll_errors = 0
        registers, raw_value, converted, unit = result
        H_m = converted * UNIT_TO_M[unit]

        self.mb_raw_label.setText(f"Register mentah: {registers}\nNilai decode: {raw_value:g}")
        self.H_input.blockSignals(True)
        self.H_input.setValue(converted)
        self.H_input.blockSignals(False)
        self.H_unit_combo.blockSignals(True)
        self.H_unit_combo.setCurrentText(unit)
        self.H_unit_combo.blockSignals(False)
        self.refresh_schematic()
        self.mb_live_h_label.setText(f"{converted:,.3f} {unit}")
        now_str = datetime.now().strftime("%H:%M:%S")
        self.mb_live_time_label.setText(f"Terakhir dibaca: {now_str} (polling aktif)")

        # Setiap polling langsung menghitung ulang debit (bukan opsional) -
        # inilah tujuan sebenarnya dari pembacaan Modbus: mengisi H dan
        # memperbarui hasil kalkulasi di tab Kalkulasi secara otomatis.
        Q_val = area_val = None
        notes_val, warnings_val, rho_val = [], [], self.rho_input.value()
        calc = self._compute_flow_quiet(H_m)
        if calc is not None:
            Q_val, area_val, notes_val, warnings_val, rho_val = calc
            self._display_results(Q_val, area_val, rho_val, notes_val, warnings_val)

        # Simpan sebagian hasil polling ke Riwayat/Export Kalkulasi, dibatasi
        # dengan interval sampling (bukan setiap siklus) supaya riwayat tidak
        # dibanjiri saat interval polling sangat cepat (bisa serendah 200ms).
        if Q_val is not None and self.mb_history_checkbox.isChecked():
            sample_interval_s = self.mb_history_interval_input.value()
            now = datetime.now()
            elapsed = (
                (now - self.last_history_save_ts).total_seconds()
                if self.last_history_save_ts is not None else None
            )
            if elapsed is None or elapsed >= sample_interval_s:
                params = self._get_param_values()
                self._append_history(
                    self.type_combo.currentText(), self.method_combo.currentText(),
                    params, converted, unit, H_m, rho_val,
                    Q_val, area_val, notes_val, warnings_val,
                    sumber="Polling Otomatis",
                )
                self.last_history_save_ts = now

        if self.mb_log_checkbox.isChecked() and self.log_csv_writer is not None:
            row = {
                "waktu": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "raw_registers": str(registers),
                "nilai_decode_mentah": raw_value,
                "H_konversi": round(converted, 4),
                "H_satuan": unit,
                "H_meter": round(H_m, 6),
                "tipe_alat_ukur": self.type_combo.currentText(),
                "metode": self.method_combo.currentText(),
                "Q_m3_per_s": round(Q_val, 6) if Q_val is not None else "",
                "Q_L_per_s": round(Q_val * 1000.0, 4) if Q_val is not None else "",
                "mdot_kg_per_s": round(rho_val * Q_val, 4) if Q_val is not None else "",
                "kecepatan_m_per_s": round(Q_val / area_val, 5) if (Q_val is not None and area_val) else "",
            }
            self._write_log_row(row)

    # ------------------------------------------------------------------
    # Dynamic form logic
    # ------------------------------------------------------------------
    def on_type_changed(self, type_name):
        if not type_name:
            return
        self.method_combo.blockSignals(True)
        self.method_combo.clear()
        self.method_combo.addItems(list(TYPES[type_name]["methods"].keys()))
        self.method_combo.blockSignals(False)
        self.on_method_changed(self.method_combo.currentText())

    def on_method_changed(self, method_name):
        type_name = self.type_combo.currentText()
        if not type_name or not method_name:
            return
        entry = TYPES[type_name]["methods"][method_name]
        self.desc_label.setText(entry.get("desc", ""))
        self._rebuild_param_form(entry["params"])
        self.refresh_schematic()

    def _rebuild_param_form(self, params):
        while self.param_form.rowCount() > 0:
            self.param_form.removeRow(0)
        self.current_param_widgets = {}

        for spec in params:
            widget = self._make_param_widget(spec)
            label_text = spec["label"]
            if spec.get("unit") and spec["unit"] != "-":
                label_text += f" [{spec['unit']}]"
            self.param_form.addRow(label_text + ":", widget)
            self.current_param_widgets[spec["key"]] = (widget, spec)

            if isinstance(widget, QComboBox):
                widget.currentTextChanged.connect(self.refresh_schematic)
            else:
                widget.valueChanged.connect(self.refresh_schematic)

    def _make_param_widget(self, spec):
        if spec["type"] == "combo":
            w = QComboBox()
            w.addItems(spec["options"])
            w.setCurrentIndex(spec.get("default_index", 0))
            return w
        else:
            w = QDoubleSpinBox()
            w.setRange(spec["minv"], spec["maxv"])
            w.setDecimals(spec["decimals"])
            w.setSingleStep(spec["step"])
            w.setValue(spec["default"])
            return w

    def _get_param_values(self):
        values = {}
        for key, (widget, spec) in self.current_param_widgets.items():
            if isinstance(widget, QComboBox):
                values[key] = widget.currentText()
            else:
                values[key] = widget.value()
        return values

    def refresh_schematic(self):
        type_name = self.type_combo.currentText()
        method_name = self.method_combo.currentText()
        if not type_name or not method_name:
            return
        params = self._get_param_values()
        H_val = self.H_input.value()
        H_unit = self.H_unit_combo.currentText()
        self.schematic.update_view(type_name, method_name, params, H_val, H_unit)

    # ------------------------------------------------------------------
    # Kalkulasi
    # ------------------------------------------------------------------
    def calculate(self, sumber="Manual (Kalkulasi)"):
        try:
            type_name = self.type_combo.currentText()
            method_name = self.method_combo.currentText()
            entry = TYPES[type_name]["methods"][method_name]
            func = entry["func"]

            H_val = self.H_input.value()
            H_unit = self.H_unit_combo.currentText()
            H_m = H_val * UNIT_TO_M[H_unit]

            if H_m <= 0:
                raise ValueError("Ketinggian air (H) harus lebih besar dari 0.")

            rho = self.rho_input.value()
            params = self._get_param_values()

            Q, area, notes, warnings = func(H_m, params, rho)

            if Q is None:
                msg = " ".join(warnings) if warnings else "Perhitungan tidak dapat diselesaikan dengan parameter ini."
                raise ValueError(msg)
            if math.isnan(Q) or math.isinf(Q) or Q < 0:
                raise ValueError("Hasil perhitungan tidak valid. Periksa kembali parameter input.")

            self._display_results(Q, area, rho, notes, warnings)
            self._append_history(type_name, method_name, params, H_val, H_unit, H_m, rho,
                                  Q, area, notes, warnings, sumber=sumber)
            self.status_label.setText(
                f"Perhitungan berhasil - {type_name} / {method_name} - "
                f"{datetime.now().strftime('%H:%M:%S')}"
            )
        except Exception as e:
            QMessageBox.warning(self, "Input Tidak Valid", str(e))
            self.status_label.setText("Terjadi kesalahan input. Periksa parameter Anda.")

    def _display_results(self, Q, area, rho, notes, warnings):
        Q_m3s = Q
        Q_m3h = Q * 3600.0
        Q_Ls = Q * 1000.0
        Q_Lmin = Q * 1000.0 * 60.0

        mdot_kgs = rho * Q
        mdot_kgh = mdot_kgs * 3600.0
        mdot_th = mdot_kgh / 1000.0

        self.val_Q.setText(f"{Q_m3s:,.5f} m\u00b3/s")
        self.sub_Q.setText(f"{Q_m3h:,.2f} m\u00b3/jam   |   {Q_Ls:,.3f} L/s   |   {Q_Lmin:,.1f} L/menit")

        self.val_M.setText(f"{mdot_kgs:,.3f} kg/s")
        self.sub_M.setText(f"{mdot_kgh:,.1f} kg/jam   |   {mdot_th:,.3f} ton/jam")

        velocity = None
        if area is not None and area > 0:
            velocity = Q / area
            self.val_V.setText(f"{velocity:,.4f} m/s")
            self.sub_V.setText("Estimasi v = Q / A pada seksi kontrol")
            self.val_A.setText(f"{area:,.5f} m\u00b2")
            self.sub_A.setText("Estimasi luas basah pada H terukur")
        else:
            self.val_V.setText("N/A")
            self.sub_V.setText("Tidak dapat diestimasi untuk metode/parameter ini")
            self.val_A.setText("N/A")
            self.sub_A.setText("Tidak dapat diestimasi untuk metode/parameter ini")

        note_lines = [f"\u2022 {n}" for n in notes]
        note_lines += [f"\u26a0 {wtext}" for wtext in warnings]
        if not note_lines:
            note_lines.append("Tidak ada catatan tambahan.")
        self.notes_label.setText("\n".join(note_lines))
        if warnings:
            self.notes_label.setStyleSheet("color:#a6410c; font-size:9pt;")
        else:
            self.notes_label.setStyleSheet("color:#445261; font-size:9pt;")

        self.last_result = dict(Q_m3s=Q_m3s, Q_m3h=Q_m3h, Q_Ls=Q_Ls, Q_Lmin=Q_Lmin,
                                 mdot_kgs=mdot_kgs, mdot_kgh=mdot_kgh, mdot_th=mdot_th,
                                 velocity=velocity, area=area)

    def _append_history(self, type_name, method_name, params, H_val, H_unit, H_m, rho,
                         Q, area, notes, warnings, sumber="Manual (Kalkulasi)"):
        r = self.last_result
        param_str = "; ".join(f"{k}={v}" for k, v in params.items())
        entry = {
            "waktu": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "sumber": sumber,
            "tipe": type_name,
            "metode": method_name,
            "parameter": param_str,
            "H_input": H_val,
            "H_satuan": H_unit,
            "H_meter": round(H_m, 6),
            "rho_kg_per_m3": rho,
            "Q_m3_per_s": round(r["Q_m3s"], 6),
            "Q_m3_per_jam": round(r["Q_m3h"], 4),
            "Q_L_per_s": round(r["Q_Ls"], 4),
            "Q_L_per_menit": round(r["Q_Lmin"], 2),
            "mdot_kg_per_s": round(r["mdot_kgs"], 4),
            "mdot_kg_per_jam": round(r["mdot_kgh"], 2),
            "mdot_ton_per_jam": round(r["mdot_th"], 4),
            "kecepatan_m_per_s": round(r["velocity"], 5) if r["velocity"] is not None else "",
            "luas_basah_m2": round(r["area"], 6) if r["area"] is not None else "",
            "catatan": " | ".join(notes) if notes else "",
            "peringatan": " | ".join(warnings) if warnings else "",
        }
        self.history.append(entry)
        self._update_history_label()

    def _update_history_label(self):
        n = len(self.history)
        if n == 0:
            self.history_label.setText("Riwayat: 0 perhitungan tersimpan di sesi ini.")
        else:
            self.history_label.setText(f"Riwayat: {n} perhitungan tersimpan di sesi ini (siap diexport).")

    # ------------------------------------------------------------------
    # Export
    # ------------------------------------------------------------------
    def export_history(self):
        if not self.history:
            QMessageBox.information(self, "Export Riwayat",
                                     "Belum ada hasil perhitungan untuk diexport.\n"
                                     "Lakukan minimal satu perhitungan (klik Hitung) terlebih dahulu.")
            return

        default_name = f"hasil_openchannel_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        fname, selected_filter = QFileDialog.getSaveFileName(
            self, "Export Riwayat Perhitungan", default_name,
            "Excel Workbook (*.xlsx);;CSV File (*.csv)"
        )
        if not fname:
            return

        try:
            if fname.lower().endswith(".xlsx") or "Excel" in selected_filter:
                if not fname.lower().endswith(".xlsx"):
                    fname += ".xlsx"
                if not OPENPYXL_AVAILABLE:
                    QMessageBox.warning(
                        self, "Modul Tidak Tersedia",
                        "Export ke Excel (.xlsx) memerlukan modul 'openpyxl' yang belum "
                        "terpasang.\n\nInstall dengan:\n    pip install openpyxl\n\n"
                        "Atau pilih format CSV sebagai alternatif."
                    )
                    return
                self._export_xlsx(fname)
            else:
                if not fname.lower().endswith(".csv"):
                    fname += ".csv"
                self._export_csv(fname)

            QMessageBox.information(self, "Export Berhasil",
                                     f"Riwayat perhitungan ({len(self.history)} baris) "
                                     f"berhasil disimpan ke:\n{fname}")
            self.status_label.setText(f"Export berhasil: {fname}")
        except Exception as e:
            QMessageBox.critical(self, "Export Gagal", f"Terjadi kesalahan saat export:\n{e}")

    def _export_csv(self, fname):
        with open(fname, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=self.HISTORY_FIELDS)
            writer.writeheader()
            for row in self.history:
                writer.writerow(row)

    def _export_xlsx(self, fname):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Riwayat Perhitungan"

        header_fill = PatternFill(start_color="14477A", end_color="14477A", fill_type="solid")
        header_font = XLFont(color="FFFFFF", bold=True)

        pretty_headers = {
            "waktu": "Waktu", "sumber": "Sumber Data", "tipe": "Tipe Alat Ukur", "metode": "Metode",
            "parameter": "Parameter", "H_input": "H (input)", "H_satuan": "Satuan H",
            "H_meter": "H (m)", "rho_kg_per_m3": "\u03c1 (kg/m3)",
            "Q_m3_per_s": "Q (m3/s)", "Q_m3_per_jam": "Q (m3/jam)",
            "Q_L_per_s": "Q (L/s)", "Q_L_per_menit": "Q (L/menit)",
            "mdot_kg_per_s": "\u1e41 (kg/s)", "mdot_kg_per_jam": "\u1e41 (kg/jam)",
            "mdot_ton_per_jam": "\u1e41 (ton/jam)", "kecepatan_m_per_s": "v (m/s)",
            "luas_basah_m2": "A (m2)", "catatan": "Catatan", "peringatan": "Peringatan",
        }

        for col_idx, key in enumerate(self.HISTORY_FIELDS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=pretty_headers.get(key, key))
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_idx, entry in enumerate(self.history, start=2):
            for col_idx, key in enumerate(self.HISTORY_FIELDS, start=1):
                ws.cell(row=row_idx, column=col_idx, value=entry.get(key, ""))

        # lebar kolom otomatis (perkiraan sederhana)
        for col_idx, key in enumerate(self.HISTORY_FIELDS, start=1):
            max_len = max(
                [len(str(pretty_headers.get(key, key)))] +
                [len(str(e.get(key, ""))) for e in self.history]
            )
            ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else "A"].width = min(max(max_len + 2, 10), 40)

        ws.freeze_panes = "A2"
        wb.save(fname)

    def clear_history(self):
        if not self.history:
            QMessageBox.information(self, "Hapus Riwayat", "Riwayat sudah kosong.")
            return
        reply = QMessageBox.question(
            self, "Hapus Riwayat",
            f"Hapus seluruh {len(self.history)} riwayat perhitungan pada sesi ini?\n"
            "Data yang belum diexport akan hilang.",
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        if reply == QMessageBox.Yes:
            self.history = []
            self._update_history_label()
            self.status_label.setText("Riwayat perhitungan telah dihapus.")

    # ------------------------------------------------------------------
    def reset_form(self):
        self.H_input.setValue(150.0)
        self.H_unit_combo.setCurrentText("mm")
        self.rho_input.setValue(1000.0)
        self.on_method_changed(self.method_combo.currentText())
        self.val_Q.setText("-")
        self.sub_Q.setText("")
        self.val_M.setText("-")
        self.sub_M.setText("")
        self.val_V.setText("-")
        self.sub_V.setText("")
        self.val_A.setText("-")
        self.sub_A.setText("")
        self.notes_label.setText("Belum ada perhitungan.")
        self.notes_label.setStyleSheet("color:#445261; font-size:9pt;")
        self.last_result = None
        self.status_label.setText("Form direset. Siap untuk perhitungan baru.")

    def show_about(self):
        xlsx_note = "tersedia" if OPENPYXL_AVAILABLE else "TIDAK tersedia (install: pip install openpyxl)"
        QMessageBox.information(
            self,
            "Tentang Aplikasi",
            "Kalkulator Open Channel Flowmeter - Professional Edition\n\n"
            "Mendukung perhitungan debit, laju alir massa, dan estimasi kecepatan "
            "aliran untuk V-Notch Weir, Rectangular Weir, Trapezoidal Weir, dan "
            "Parshall Flume dengan beberapa metode/formula empiris standar.\n\n"
            f"Export ke Excel (.xlsx): {xlsx_note}\n"
            "Export ke CSV: selalu tersedia.\n\n"
            "Catatan: Formula & koefisien default bersifat umum/tipikal. Untuk "
            "pengukuran resmi atau legal metering, verifikasi terhadap standar "
            "yang berlaku (ISO 1438, ISO 9826, SNI, atau kalibrasi lapangan)."
        )

    def closeEvent(self, event):
        """Membersihkan koneksi Modbus, timer polling, dan file log saat aplikasi ditutup."""
        try:
            if self.is_polling:
                self.poll_timer.stop()
            self._close_log_file()
            if self.modbus_reader.connected:
                self.modbus_reader.disconnect()
        except Exception:
            pass
        event.accept()


def main():
    app = QApplication(sys.argv)
    app.setStyleSheet(APP_STYLE)
    win = MainWindow()
    win.show()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()
